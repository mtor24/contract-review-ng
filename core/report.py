"""Review report export: assemble report data from a stored analysis and
render it as PDF (reportlab) or DOCX (python-docx).

All contract text is untrusted input; every value that reaches a reportlab
Paragraph passes through xml.sax.saxutils.escape so it can never become
markup.
"""

import io
from datetime import date
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from core.classify import LABEL_NAMES
from core.missing import contract_types

DISCLAIMER = (
    "This report was produced by a tool that assists contract review. "
    "It does not give legal advice. Every finding is a suggestion until "
    "a lawyer accepts it."
)

_FOOTER_DISCLAIMER = "Assists contract review. Does not give legal advice."

NAVY = colors.HexColor("#1E2A44")
GOLD = colors.HexColor("#B8963E")
HEADER_BG = colors.HexColor("#F7F8FA")
GRID = colors.HexColor("#E5E7EB")

_SEVERITY_COLOURS = {
    "High": colors.HexColor("#C0392B"),
    "Medium": colors.HexColor("#D68910"),
    "Low": colors.HexColor("#2E86C1"),
    "Pass": colors.HexColor("#1E8449"),
}

_STATUS_LABELS = {
    "suggested": "Suggestion",
    "accepted": "Accepted",
    "rejected": "Rejected",
    "edited": "Edited",
}

# Lawyer-decisions section only lists findings a lawyer has acted on.
_DECIDED_STATUSES = {"accepted", "rejected", "edited"}

# Finding kind shown in words in the decisions table.
_KIND_NAMES = {
    "risk": "Risk",
    "missing": "Missing clause",
    "compliance": "Compliance",
    "clause": "Clause",
    "obligation": "Obligation",
}


def build_report_data(
    contract: dict, version: dict, summary: dict, firm_name: str
) -> dict:
    """Assemble everything a report needs from stored shapes."""
    parties = "; ".join(
        p["name"] + (f' ("{p["role"]}")' if p.get("role") else "")
        for p in contract.get("parties") or []
    )
    clauses = [
        {
            "number": c.get("number") or "",
            "heading": c.get("heading") or "",
            "label_name": LABEL_NAMES.get(c.get("label"), c.get("label") or ""),
        }
        for c in version["clauses"]
    ]
    findings = version["findings"]
    missing = [f for f in findings if f["kind"] == "missing"]
    risks = [f for f in findings if f["kind"] == "risk"]
    compliance = [f for f in findings if f["kind"] == "compliance"]
    obligations = [
        {
            "party": o["party"],
            "period": o.get("period_text") or "",
            "due_date": o.get("due_date") or "",
            "excerpt": o.get("text") or "",
        }
        for o in version["obligations"]
    ]
    decisions = [
        {
            "kind": _KIND_NAMES.get(f["kind"], f["kind"]),
            "name": f.get("name") or f.get("label") or "",
            "severity": f["severity"],
            "status": f["status"],
            "note": f.get("note") or "",
            "edited_text": f.get("edited_text") or "",
        }
        for f in findings
        if f["status"] in _DECIDED_STATUSES
    ]
    total = len(findings)
    reviewed = sum(1 for f in findings if f["status"] != "suggested")
    return {
        "cover": {
            "firm_name": firm_name,
            "title": contract["title"],
            # Human type name ("Supply Agreement"), never the id.
            "contract_type": contract_types().get(
                contract["type"], contract["type"]
            ),
            "parties": parties,
            "version_no": version["version_no"],
            "generated": date.today().isoformat(),
            "value_text": summary.get("value_text") or "",
        },
        "summary": {
            "agreement_date": summary.get("agreement_date") or "",
            "effective_date": summary.get("effective_date") or "",
            "expiry_date": summary.get("expiry_date") or "",
            "value_text": summary.get("value_text") or "",
            "term_text": summary.get("term_text") or "",
            "key_terms": summary.get("key_terms") or [],
            "counts": summary.get("counts") or {},
        },
        "clauses": clauses,
        "missing": missing,
        "risks": risks,
        "compliance": compliance,
        "obligations": obligations,
        "decisions": decisions,
        "decision_counts": {
            "reviewed": reviewed,
            "total": total,
            "awaiting": total - reviewed,
        },
        "disclaimer": DISCLAIMER,
    }


# ---------------------------------------------------------------- PDF

# WHY keepWithNext=1: a section heading must never sit alone at the
# bottom of a page; reportlab keeps it with the flowable that follows.
_HEADING_STYLE = ParagraphStyle(
    "SectionHeading", fontName="Helvetica-Bold", fontSize=13, textColor=NAVY,
    spaceBefore=14, spaceAfter=6, keepWithNext=1,
)
_FIRM_STYLE = ParagraphStyle(
    "FirmName", fontName="Helvetica-Bold", fontSize=16, textColor=NAVY,
    spaceAfter=6,
)
_BAND_KICKER_STYLE = ParagraphStyle(
    "BandKicker", fontName="Helvetica-Bold", fontSize=9, textColor=GOLD,
)
_BAND_TITLE_STYLE = ParagraphStyle(
    "BandTitle", fontName="Helvetica-Bold", fontSize=18, leading=22,
    textColor=colors.white,
)
_BODY_STYLE = ParagraphStyle(
    "Body", fontName="Helvetica", fontSize=9, leading=12,
)
_SMALL_STYLE = ParagraphStyle(
    "Small", fontName="Helvetica", fontSize=8, leading=10,
)
_DISCLAIMER_STYLE = ParagraphStyle(
    "Disclaimer", fontName="Helvetica-Oblique", fontSize=9, leading=12,
    textColor=colors.HexColor("#444444"),
)


def _p(text: str, style: ParagraphStyle = _BODY_STYLE) -> Paragraph:
    # Contract text is untrusted: escape before it becomes Paragraph markup.
    return Paragraph(escape(str(text)), style)


def _severity_paragraph(severity: str) -> Paragraph:
    colour = _SEVERITY_COLOURS.get(severity, colors.black)
    style = ParagraphStyle(
        f"sev-{severity}", parent=_BODY_STYLE, textColor=colour,
        fontName="Helvetica-Bold",
    )
    return Paragraph(escape(severity), style)


def _table(header: list[str], rows: list[list]) -> Table:
    data = [[_p(h, ParagraphStyle("th", parent=_BODY_STYLE, fontName="Helvetica-Bold")) for h in header]]
    data.extend(rows)
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
                ("GRID", (0, 0), (-1, -1), 0.5, GRID),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _title_band(title: str) -> Table:
    # Navy full-width band: small gold kicker above the large white title.
    band = Table(
        [
            [_p("Contract Review Report", _BAND_KICKER_STYLE)],
            [_p(title, _BAND_TITLE_STYLE)],
        ],
        colWidths=[A4[0] - 36 * mm],
    )
    band.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), NAVY),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
            ]
        )
    )
    return band


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(18 * mm, 10 * mm, _FOOTER_DISCLAIMER)
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def to_pdf(data: dict) -> bytes:
    """Render the report as a PDF; returns raw bytes."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    )
    story = []

    # Cover: navy title band, firm name below it.
    cover = data["cover"]
    story.append(_title_band(cover["title"]))
    story.append(Spacer(1, 6))
    story.append(_p(cover["firm_name"] or "Firm", _FIRM_STYLE))
    story.append(Spacer(1, 4))
    cover_rows = [
        ["Contract type", cover["contract_type"]],
        ["Parties", cover["parties"]],
        ["Version", str(cover["version_no"])],
        ["Generated", cover["generated"]],
    ]
    if cover["value_text"]:
        cover_rows.append(["Contract value", cover["value_text"]])
    story.append(_table(
        ["Field", "Value"],
        [[_p(k), _p(v)] for k, v in cover_rows],
    ))
    story.append(Spacer(1, 10))
    # The disclaimer is boxed so it cannot be missed.
    disclaimer_box = Table([[_p(data["disclaimer"], _DISCLAIMER_STYLE)]])
    disclaimer_box.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.75, NAVY),
                ("BACKGROUND", (0, 0), (-1, -1), HEADER_BG),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(disclaimer_box)

    # Summary
    summary = data["summary"]
    story.append(_p("Summary", _HEADING_STYLE))
    summary_rows = []
    for label, key in (
        ("Agreement date", "agreement_date"),
        ("Effective date", "effective_date"),
        ("Expiry date", "expiry_date"),
        ("Contract value", "value_text"),
        ("Term", "term_text"),
    ):
        if summary.get(key):
            summary_rows.append([_p(label), _p(summary[key])])
    counts = summary.get("counts") or {}
    if counts:
        summary_rows.append(
            [_p("Clauses identified"), _p(str(counts.get("clauses", 0)))]
        )
        summary_rows.append(
            [
                _p("Open risks (High/Medium/Low)"),
                _p(
                    f"{counts.get('risks_high', 0)} / "
                    f"{counts.get('risks_medium', 0)} / "
                    f"{counts.get('risks_low', 0)}"
                ),
            ]
        )
    if summary_rows:
        story.append(_table(["Field", "Value"], summary_rows))
    if summary.get("key_terms"):
        story.append(Spacer(1, 6))
        story.append(_table(
            ["Key term", "Sentence"],
            [[_p(t["name"]), _p(t["sentence"])] for t in summary["key_terms"]],
        ))

    # Clauses identified
    story.append(_p("Clauses identified", _HEADING_STYLE))
    if data["clauses"]:
        story.append(_table(
            ["No.", "Heading", "Type"],
            [
                [_p(c["number"]), _p(c["heading"]), _p(c["label_name"])]
                for c in data["clauses"]
            ],
        ))
    else:
        story.append(_p("No clauses identified."))

    # Missing clauses: show the human name ("Notices"), not the id.
    story.append(_p("Missing clauses", _HEADING_STYLE))
    if data["missing"]:
        story.append(_table(
            ["Clause", "Severity", "Reason"],
            [
                [
                    _p(f.get("name") or f.get("label") or ""),
                    _severity_paragraph(f["severity"]),
                    _p(f["reason"]),
                ]
                for f in data["missing"]
            ],
        ))
    else:
        story.append(_p("No missing clauses detected."))

    # Risks
    story.append(_p("Risks", _HEADING_STYLE))
    if data["risks"]:
        story.append(_table(
            ["Severity", "Risk", "Why it matters", "Source text", "Status"],
            [
                [
                    _severity_paragraph(f["severity"]),
                    _p(f.get("name") or f.get("label") or ""),
                    _p(f["reason"]),
                    _p(f.get("excerpt") or "", _SMALL_STYLE),
                    _p(_STATUS_LABELS.get(f.get("status", "suggested"), f.get("status", ""))),
                ]
                for f in data["risks"]
            ],
        ))
    else:
        story.append(_p("No risks detected."))

    # Nigerian compliance (indicative)
    story.append(_p("Nigerian compliance (indicative)", _HEADING_STYLE))
    if data["compliance"]:
        story.append(_table(
            ["Check", "Result", "Detail"],
            [
                [
                    _p(f.get("name") or f.get("label") or ""),
                    _severity_paragraph(f["severity"]),
                    _p(f["reason"]),
                ]
                for f in data["compliance"]
            ],
        ))
    else:
        story.append(_p("No compliance checks applied."))

    # Obligations and deadlines
    story.append(_p("Obligations and deadlines", _HEADING_STYLE))
    if data["obligations"]:
        story.append(_table(
            ["Party", "Period", "Due date", "Obligation"],
            [
                [
                    _p(o["party"]),
                    _p(o["period"]),
                    _p(o["due_date"]),
                    _p(o["excerpt"], _SMALL_STYLE),
                ]
                for o in data["obligations"]
            ],
        ))
    else:
        story.append(_p("No obligations detected."))

    # Lawyer decisions: only findings a lawyer has accepted, rejected or edited.
    story.append(_p("Lawyer decisions", _HEADING_STYLE))
    counts = data.get("decision_counts") or {}
    if counts:
        story.append(_p(
            f"{counts['reviewed']} of {counts['total']} findings reviewed. "
            f"{counts['awaiting']} still awaiting review."
        ))
        story.append(Spacer(1, 4))
    if data["decisions"]:
        rows = []
        for d in data["decisions"]:
            status = _STATUS_LABELS.get(d["status"], d["status"])
            note = d["note"]
            if d["edited_text"]:
                note = (note + " " if note else "") + f'Edited: {d["edited_text"]}'
            rows.append(
                [
                    _p(d["kind"]),
                    _p(d["name"]),
                    _severity_paragraph(d["severity"]),
                    _p(status),
                    _p(note, _SMALL_STYLE),
                ]
            )
        story.append(_table(
            ["Type", "Finding", "Severity", "Decision", "Note / edited text"],
            rows,
        ))
    else:
        story.append(_p("No decisions recorded yet."))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# ---------------------------------------------------------------- DOCX


def to_docx(data: dict) -> bytes:
    """Render the report as a DOCX; returns raw bytes."""
    import docx

    document = docx.Document()

    cover = data["cover"]
    document.add_heading(cover["firm_name"] or "Firm", level=0)
    document.add_paragraph(cover["title"]).runs[0].bold = True
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (
        ("Contract type", cover["contract_type"]),
        ("Parties", cover["parties"]),
        ("Version", str(cover["version_no"])),
        ("Generated", cover["generated"]),
        ("Contract value", cover["value_text"]),
    ):
        if value == "":
            continue
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
    p = document.add_paragraph()
    p.add_run(data["disclaimer"]).italic = True

    def add_table(header, rows):
        if not rows:
            return
        t = document.add_table(rows=1, cols=len(header))
        try:
            t.style = "Light Grid Accent 1"
        except KeyError:
            # Style name varies by Word version; fall back to a plain grid.
            t.style = "Table Grid"
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)

    summary = data["summary"]
    document.add_heading("Summary", level=1)
    add_table(
        ["Field", "Value"],
        [
            (label, summary[key])
            for label, key in (
                ("Agreement date", "agreement_date"),
                ("Effective date", "effective_date"),
                ("Expiry date", "expiry_date"),
                ("Contract value", "value_text"),
                ("Term", "term_text"),
            )
            if summary.get(key)
        ],
    )
    add_table(
        ["Key term", "Sentence"],
        [(t["name"], t["sentence"]) for t in summary.get("key_terms") or []],
    )

    document.add_heading("Clauses identified", level=1)
    add_table(
        ["No.", "Heading", "Type"],
        [(c["number"], c["heading"], c["label_name"]) for c in data["clauses"]],
    )

    document.add_heading("Missing clauses", level=1)
    add_table(
        ["Clause", "Severity", "Reason"],
        [
            (f.get("name") or f.get("label") or "", f["severity"], f["reason"])
            for f in data["missing"]
        ],
    )

    document.add_heading("Risks", level=1)
    add_table(
        ["Severity", "Risk", "Why it matters", "Source text", "Status"],
        [
            (
                f["severity"],
                f.get("name") or f.get("label") or "",
                f["reason"],
                f.get("excerpt") or "",
                _STATUS_LABELS.get(f.get("status", "suggested"), f.get("status", "")),
            )
            for f in data["risks"]
        ],
    )

    document.add_heading("Nigerian compliance (indicative)", level=1)
    add_table(
        ["Check", "Result", "Detail"],
        [
            (f.get("name") or f.get("label") or "", f["severity"], f["reason"])
            for f in data["compliance"]
        ],
    )

    document.add_heading("Obligations and deadlines", level=1)
    add_table(
        ["Party", "Period", "Due date", "Obligation"],
        [
            (o["party"], o["period"], o["due_date"], o["excerpt"])
            for o in data["obligations"]
        ],
    )

    document.add_heading("Lawyer decisions", level=1)
    counts = data.get("decision_counts") or {}
    if counts:
        document.add_paragraph(
            f"{counts['reviewed']} of {counts['total']} findings reviewed. "
            f"{counts['awaiting']} still awaiting review."
        )
    decision_rows = []
    for d in data["decisions"]:
        status = _STATUS_LABELS.get(d["status"], d["status"])
        note = d["note"]
        if d["edited_text"]:
            note = (note + " " if note else "") + f'Edited: {d["edited_text"]}'
        decision_rows.append(
            (d["kind"], d["name"], d["severity"], status, note)
        )
    if decision_rows:
        add_table(
            ["Type", "Finding", "Severity", "Decision", "Note / edited text"],
            decision_rows,
        )
    else:
        document.add_paragraph("No decisions recorded yet.")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
