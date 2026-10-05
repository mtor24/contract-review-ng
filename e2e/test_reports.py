"""Reports: preview with disclaimer, export PDF and Word."""

import pytest

from e2e.conftest import expect, open_page, wait_idle

DISCLAIMER = (
    "This report was produced by a tool that assists contract review. "
    "It does not give legal advice."
)


def _open_report(page, base_url: str, title: str) -> None:
    open_page(page, base_url, "reports")
    # The selectbox opens from its combobox input, not from the label text.
    page.get_by_role("combobox", name="Contract").click()
    page.get_by_role("option", name=title).click()
    wait_idle(page)


@pytest.mark.cell("reports.export", "owner")
def test_preview_pdf_and_word_exports(page, owner_server, tmp_path):
    _open_report(page, owner_server.url, "Supply Agreement")

    # Preview shows the disclaimer.
    expect(page.get_by_text(DISCLAIMER, exact=False)).to_be_visible()

    # PDF download starts with %PDF.
    with page.expect_download() as pdf_dl:
        page.get_by_role("button", name="Export PDF").click()
    pdf_path = tmp_path / "report.pdf"
    pdf_dl.value.save_as(pdf_path)
    assert pdf_path.read_bytes().startswith(b"%PDF")
    wait_idle(page)

    # Word download starts with PK and opens with python-docx, disclaimer included.
    with page.expect_download() as docx_dl:
        page.get_by_role("button", name="Export Word").click()
    docx_path = tmp_path / "report.docx"
    docx_dl.value.save_as(docx_path)
    assert docx_path.read_bytes().startswith(b"PK")

    import docx

    document = docx.Document(str(docx_path))
    text = "\n".join(p.text for p in document.paragraphs)
    assert DISCLAIMER in text
