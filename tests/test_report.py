"""Tests for core.report: PDF and DOCX export from a real stored analysis."""

import io

import docx
import pytest

from core import pipeline
from core.report import DISCLAIMER, build_report_data, to_docx, to_pdf
from storage import db, models

FIRM_NAME = "Test & Co. Chambers"


def _docx_text(raw: bytes) -> str:
    document = docx.Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    return text


def _decisions_table_text(raw: bytes) -> str:
    # The decisions table is the one headed "Type", "Finding", ...
    document = docx.Document(io.BytesIO(raw))
    for table in document.tables:
        header = [c.text for c in table.rows[0].cells]
        if header[:2] == ["Type", "Finding"]:
            text = ""
            for row in table.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
            return text
    return ""


@pytest.fixture
def stored(tmp_path, samples_dir):
    """A real analysed supply agreement stored in a temp database, with one
    finding accepted."""
    text = (samples_dir / "supply_agreement.txt").read_text(encoding="utf-8")
    result = pipeline.analyse(text)
    conn = db.connect(tmp_path / "test.db")
    contract_id, version_id = models.save_analysis(
        conn, result, "supply_agreement.txt"
    )
    # Accept the first finding so the report has one real lawyer decision.
    first_fid = result["findings"][0]["fid"]
    models.set_decision(conn, version_id, first_fid, "accepted", note="Looks fine")

    contract = models.get_contract(conn, contract_id)
    version = models.get_version(conn, version_id)
    data = build_report_data(contract, version, result["summary"], FIRM_NAME)
    conn.close()
    return data


def test_to_pdf_returns_valid_pdf(stored):
    pdf = to_pdf(stored)
    assert isinstance(pdf, bytes)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 2000


def test_to_docx_roundtrip_contains_key_text(stored):
    raw = to_docx(stored)
    assert isinstance(raw, bytes)
    text = _docx_text(raw)
    assert DISCLAIMER in text
    assert FIRM_NAME in text
    assert "Accepted" in text


def test_docx_decisions_lists_only_decided_findings(stored):
    raw = to_docx(stored)
    decisions_text = _decisions_table_text(raw)
    assert "Accepted" in decisions_text
    # Only one finding was decided: the table must not list every obligation.
    assert decisions_text.lower().count("obligation") < 5


def test_missing_clause_shows_human_name(stored):
    # The supply agreement lacks a notices clause; the report must show
    # the human name "Notices", not the id "notices".
    assert "Notices" in _docx_text(to_docx(stored))
    assert "Notices" in [f["name"] for f in stored["missing"]]
