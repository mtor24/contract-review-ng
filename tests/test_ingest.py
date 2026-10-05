import docx
import pytest
from reportlab.pdfgen import canvas

from core.ingest import clean_text, extract_text


def test_clean_text_replaces_curly_quotes_and_dashes():
    raw = "\u2018hello\u2019 \u201cworld\u201d \u2014 dash \u2013 end"
    assert clean_text(raw) == "'hello' \"world\" - dash - end"


def test_clean_text_removes_page_lines():
    raw = "First line\nPage 2 of 5\nSecond line"
    assert clean_text(raw) == "First line\nSecond line"


def test_clean_text_joins_hyphenated_words():
    assert clean_text("termi-\nnation") == "termination"


def test_clean_text_collapses_blank_lines():
    assert clean_text("a\n\n\n\nb") == "a\n\nb"


def test_extract_text_txt_bytes():
    assert extract_text("hello world".encode("utf-8"), "note.txt") == "hello world"


def test_extract_text_docx(tmp_path):
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell one"
    table.cell(0, 1).text = "Cell two"
    document.save(path)

    text = extract_text(path.read_bytes(), "doc.docx")
    for expected in ("First paragraph", "Second paragraph", "Cell one", "Cell two"):
        assert expected in text


def test_extract_text_pdf(tmp_path):
    path = tmp_path / "doc.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(100, 750, "Alpha line")
    c.drawString(100, 730, "Beta line")
    c.save()

    text = extract_text(path.read_bytes(), "doc.pdf")
    assert "Alpha line" in text
    assert "Beta line" in text


def test_extract_text_unsupported_type():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(b"data", "sheet.xls")


def _numbered_docx(tmp_path, paragraphs):
    """Build a .docx from (style, text) pairs; style None means plain text."""
    path = tmp_path / "numbered.docx"
    document = docx.Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    document.save(path)
    return extract_text(path.read_bytes(), "numbered.docx")


def test_extract_text_docx_rebuilds_word_numbering(tmp_path):
    # python-docx drops numbers that Word generates, so ingest rebuilds them.
    text = _numbered_docx(
        tmp_path,
        [
            (None, "SUPPLY AGREEMENT"),
            ("List Number", "DEFINITIONS"),
            ("List Number 2", "Words have these meanings."),
            ("List Number 2", "Headings do not affect meaning."),
            ("List Number", "PAYMENT"),
            ("List Number 2", "The Buyer shall pay within 30 days."),
            ("List Number 3", "by bank transfer;"),
            ("List Number 3", "in Naira."),
            ("List Number 2", "Interest runs on late sums."),
        ],
    )
    assert text.split("\n") == [
        "SUPPLY AGREEMENT",
        "1. DEFINITIONS",
        "1.1 Words have these meanings.",
        "1.2 Headings do not affect meaning.",
        "2. PAYMENT",
        # The child counter restarts when the parent level advances.
        "2.1 The Buyer shall pay within 30 days.",
        "(a) by bank transfer;",
        "(b) in Naira.",
        # Deeper "(a)" counters restart when level 1 advances; level 1 continues.
        "2.2 Interest runs on late sums.",
    ]


def test_extract_text_docx_ignores_bullets_and_plain_paragraphs(tmp_path):
    text = _numbered_docx(
        tmp_path,
        [
            (None, "Plain paragraph"),
            ("List Bullet", "A bullet point"),
            ("List Number", "First numbered"),
        ],
    )
    assert text.split("\n") == ["Plain paragraph", "A bullet point", "1. First numbered"]


def test_docx_numbering_lets_segmenter_find_clauses(tmp_path):
    from core.segment import segment

    text = _numbered_docx(
        tmp_path,
        [
            (None, "SUPPLY AGREEMENT"),
            ("List Number", "DEFINITIONS"),
            ("List Number 2", "Words have these meanings."),
            ("List Number", "GOVERNING LAW"),
            ("List Number 2", "Nigerian law applies."),
        ],
    )
    numbers = [c.number for c in segment(text) if c.number]
    assert numbers == ["1", "1.1", "2", "2.1"]


def test_extract_text_docx_direct_numbering_uses_ilvl(tmp_path):
    # Real multi-level lists put w:numPr (numId + ilvl) on the paragraph itself.
    path = tmp_path / "direct.docx"
    document = docx.Document()
    for level, text in [(0, "TERM"), (1, "One year."), (1, "Renewal."), (0, "NOTICES")]:
        paragraph = document.add_paragraph(text)
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_numId().val = 5  # a decimal list in the default template
        num_pr.get_or_add_ilvl().val = level
    document.save(path)

    text = extract_text(path.read_bytes(), "direct.docx")
    assert text.split("\n") == ["1. TERM", "1.1 One year.", "1.2 Renewal.", "2. NOTICES"]
